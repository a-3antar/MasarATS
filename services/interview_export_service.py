"""تصدير أسئلة الوظيفة إلى ملف Word قابل للتعبئة، واستيراد الإجابات منه بعد التعبئة.

كل سؤال يحمل معرّفاً مخفياً [Q#<id>] في نفس فقرة السؤال، ويحمل الملف كله معرّف الوظيفة
[JOB#<id>]. المطابقة عند الاستيراد تتم بالمعرّف فقط (وليس بنص السؤال) فلا تتأثر بأي تعديل
عرضي في صياغة السؤال، ولا يُقبل ملف تابع لوظيفة أخرى.
"""

import io
import re
from typing import TYPE_CHECKING

from core.exceptions import DocumentParsingError, ValidationError
from core.logging import get_logger

if TYPE_CHECKING:  # للـ type hints فقط - بدون استيراد فعلي
    from models.job import Job
    from models.job_question import JobQuestion

logger = get_logger(__name__)

_QUESTION_MARKER_RE = re.compile(r"\[Q#(\d+)\]")
_JOB_MARKER_RE = re.compile(r"\[JOB#(\d+)\]")
_ANSWER_LABEL = "الإجابة:"
_ANSWER_LABEL_RE = re.compile(r"^الإجابة\s*[:：]")
_SEPARATOR_CHAR = "─"
_SEPARATOR = _SEPARATOR_CHAR * 40
_ANSWER_BLANK_LINES = 3  # عدد الأسطر الفارغة تحت كل سؤال كمساحة كتابة
_RATIONALE_PREFIX = "💡"

_CATEGORY_LABELS = {
    "cv_specific": "خاص بالسيرة الذاتية",
    "technical": "تقني",
    "behavioral": "سلوكي",
    "leadership": "قيادي",
}

_PPR_SUCCESSORS = (
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)


# ------------------------------------------------------------------ التصدير

def _make_rtl(paragraph) -> None:
    """يجعل اتجاه الفقرة من اليمين لليسار (يُدرَج w:bidi في موضعه الصحيح حسب مخطط OOXML)."""
    from docx.oxml import OxmlElement

    properties = paragraph._p.get_or_add_pPr()
    if properties.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi") is None:
        properties.insert_element_before(OxmlElement("w:bidi"), *_PPR_SUCCESSORS)


def _add_run(paragraph, text: str, *, bold: bool = False, italic: bool = False, size: float | None = None):
    from docx.shared import Pt

    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.rtl = True
    if size:
        run.font.size = Pt(size)
    return run


def _add_hidden_marker(paragraph, marker: str) -> None:
    """ماركر شبه غير مرئي (خط 1pt أبيض) لكنه نص عادي يقرؤه python-docx عند الاستيراد."""
    from docx.shared import Pt, RGBColor

    run = paragraph.add_run(marker)
    run.font.size = Pt(1)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def export_questions_docx(job: "Job", questions: list["JobQuestion"]) -> bytes:
    """
    يبني ملف Word فيه كل أسئلة بنك الوظيفة، وتحت كل سؤال مساحة لكتابة الإجابة.
    يرفع ValidationError إذا لم توجد أسئلة.
    """
    if not questions:
        raise ValidationError("لا توجد أسئلة في بنك هذه الوظيفة لتصديرها. أضف أسئلة أولاً.")

    import docx

    document = docx.Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = docx.shared.Pt(11)

    title = document.add_heading(level=1)
    _add_hidden_marker(title, f"[JOB#{job.id}]")
    _add_run(title, f"أسئلة مقابلة: {job.title}")
    _make_rtl(title)

    meta = document.add_paragraph()
    details = " · ".join(p for p in (job.department, job.location) if p)
    _add_run(meta, (details + " · " if details else "") + f"عدد الأسئلة: {len(questions)}", italic=True)
    _make_rtl(meta)

    instructions = document.add_paragraph()
    _add_run(
        instructions,
        "تعليمات: اكتب إجابة المرشح مباشرة تحت كل سؤال في المساحة المخصصة. لا تحذف السؤال ولا "
        "تغيّر ترتيب الأسئلة أو تنقل الإجابة إلى تحت سؤال آخر. بعد الانتهاء احفظ الملف وأعد رفعه "
        "من صفحة «المقابلات» داخل بطاقة المقابلة لتصحيح الإجابات وتقييمها.",
        italic=True, size=10,
    )
    _make_rtl(instructions)
    document.add_paragraph()

    for index, question in enumerate(questions, start=1):
        heading = document.add_paragraph()
        _add_hidden_marker(heading, f"[Q#{question.id}]")
        label = _CATEGORY_LABELS.get(question.category or "", "")
        prefix = f"{index}. " + (f"[{label}] " if label else "")
        _add_run(heading, prefix + question.question, bold=True, size=12)
        _make_rtl(heading)

        if question.rationale:
            rationale = document.add_paragraph()
            _add_run(rationale, f"{_RATIONALE_PREFIX} {question.rationale}", italic=True, size=9)
            _make_rtl(rationale)

        answer_label = document.add_paragraph()
        _add_run(answer_label, _ANSWER_LABEL, bold=True)
        _make_rtl(answer_label)

        for _ in range(_ANSWER_BLANK_LINES):
            _make_rtl(document.add_paragraph())

        separator = document.add_paragraph(_SEPARATOR)
        _make_rtl(separator)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ الاستيراد

def _extract_answer(lines: list[str]) -> str:
    """
    يستخرج نص الإجابة من الأسطر الواقعة بين ماركر سؤال والماركر التالي.
    - الحالة العادية: كل ما يلي «الإجابة:» (بما فيه ما كُتب في نفس سطرها).
    - إن حذف المستخدم تسمية «الإجابة:»: نأخذ كل الأسطر ما عدا الملاحظة التوضيحية والفاصل.
    """
    label_index = next((i for i, line in enumerate(lines) if _ANSWER_LABEL_RE.match(line)), None)
    if label_index is not None:
        remainder = _ANSWER_LABEL_RE.sub("", lines[label_index], count=1).strip()
        body = ([remainder] if remainder else []) + lines[label_index + 1:]
    else:
        body = [line for line in lines if not line.startswith(_RATIONALE_PREFIX)]

    body = [line for line in body if set(line) != {_SEPARATOR_CHAR}]
    return "\n".join(body).strip()


def parse_answers_docx(file_bytes: bytes, expected_job_id: int | None = None) -> dict[int, str]:
    """
    يقرأ ملف Word معبّأ ويرجع {question_id: answer_text} للأسئلة التي كُتبت لها إجابة فعلاً.
    - expected_job_id: إن مُرِّر ووُجد معرّف وظيفة في الملف يجب أن يطابقه، وإلا ValidationError.
    - يرفع DocumentParsingError إذا لم يكن الملف صالحاً أو لا يحتوي أي علامة سؤال.
    """
    try:
        import docx
    except ImportError as exc:
        raise DocumentParsingError("مكتبة python-docx غير مثبّتة (pip install python-docx).") from exc

    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise DocumentParsingError(f"تعذّر فتح الملف كملف Word صالح: {exc}") from exc

    file_job_id: int | None = None
    sections: list[tuple[int, list[str]]] = []
    current_lines: list[str] | None = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        job_match = _JOB_MARKER_RE.search(text)
        if job_match and file_job_id is None:
            file_job_id = int(job_match.group(1))

        question_match = _QUESTION_MARKER_RE.search(text)
        if question_match:
            current_lines = []
            sections.append((int(question_match.group(1)), current_lines))
            continue

        if current_lines is not None and text:
            current_lines.append(text)

    if not sections:
        raise DocumentParsingError(
            "لم يتم العثور على أي علامة سؤال في الملف. ارفع نفس الملف المُصدَّر من البرنامج "
            "دون إعادة تصميمه أو نسخ محتواه إلى ملف جديد."
        )

    if expected_job_id is not None and file_job_id is not None and file_job_id != expected_job_id:
        raise ValidationError("هذا الملف يخص وظيفة أخرى غير الوظيفة المحددة حالياً.")

    answers: dict[int, str] = {}
    for question_id, lines in sections:
        answer = _extract_answer(lines)
        if answer:
            answers[question_id] = answer

    logger.info("Parsed %s answered question(s) from uploaded Word file", len(answers))
    return answers
