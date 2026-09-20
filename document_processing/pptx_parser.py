"""استخلاص النص والصورة من ملفات PowerPoint (.pptx) باستخدام python-pptx."""

import io

from core.constants import PHOTO_SEARCH_MAX_PAGES
from core.exceptions import DocumentParsingError
from core.logging import get_logger
from document_processing.base import DocumentParser, PhotoCandidate, pick_best_photo

logger = get_logger(__name__)


class PowerPointParser(DocumentParser):
    """يستخلص النص من الشرائح (مربعات النص، الجداول، المجموعات) ومن الصور الشخصية."""

    def extract_text(self, file_path: str) -> str:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise DocumentParsingError("مكتبة python-pptx غير مثبّتة (pip install python-pptx).") from exc

        try:
            presentation = Presentation(file_path)
            parts: list[str] = []
            for slide in presentation.slides:
                for shape in slide.shapes:
                    self._collect_text(shape, parts)

            text = "\n".join(parts).strip()
            if not text:
                raise DocumentParsingError("الملف لا يحتوي على نص قابل للاستخلاص (قد تكون الشرائح صوراً فقط).")
            return text
        except DocumentParsingError:
            raise
        except Exception as exc:
            raise DocumentParsingError(f"فشل تحليل ملف PowerPoint: {exc}") from exc

    @classmethod
    def _collect_text(cls, shape, parts: list[str]) -> None:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                cls._collect_text(child, parts)
            return

        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text.strip())

        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip(" |"):
                    parts.append(row_text)

    def extract_photo(self, file_path: str) -> tuple[bytes, str] | None:
        """يبحث عن أنسب صورة شخصية في أول شرائح الملف."""
        try:
            from PIL import Image
            from pptx import Presentation
        except ImportError:
            return None

        try:
            presentation = Presentation(file_path)
            found: list[PhotoCandidate] = []
            for slide in list(presentation.slides)[:PHOTO_SEARCH_MAX_PAGES]:
                for picture in self._iter_pictures(slide.shapes):
                    blob = picture.image.blob
                    try:
                        with Image.open(io.BytesIO(blob)) as image:
                            found.append((blob, (image.format or "png").lower(), image.width, image.height))
                    except Exception:  # noqa: BLE001 - صيغ غير مقروءة لـ Pillow نتجاهلها
                        continue
            return pick_best_photo(found)
        except Exception as exc:
            logger.warning("Photo extraction from PPTX failed: %s", exc)
            return None

    @classmethod
    def _iter_pictures(cls, shapes):
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from cls._iter_pictures(shape.shapes)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                yield shape