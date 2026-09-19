"""استخلاص النص من ملفات Word (.docx) باستخدام python-docx."""

from core.exceptions import DocumentParsingError
from document_processing.base import DocumentParser


class WordParser(DocumentParser):
    """يستخلص النص من ملفات .docx (الفقرات والجداول)."""

    def extract_text(self, file_path: str) -> str:
        try:
            import docx
        except ImportError as exc:
            raise DocumentParsingError("مكتبة python-docx غير مثبّتة (pip install python-docx).") from exc

        try:
            document = docx.Document(file_path)
            parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]

            for table in document.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip(" |"):
                        parts.append(row_text)

            text = "\n".join(parts).strip()
            if not text:
                raise DocumentParsingError("الملف لا يحتوي على نص قابل للاستخلاص.")
            return text
        except DocumentParsingError:
            raise
        except Exception as exc:
            raise DocumentParsingError(f"فشل تحليل ملف Word: {exc}") from exc
